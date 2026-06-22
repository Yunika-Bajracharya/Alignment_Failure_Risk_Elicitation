#!/usr/bin/env python3
"""Divide expert profiles into role-diverse workshop groups + build a slide.

Reads ``experts.csv`` (schema: personal_id, affiliation, role, years_experience,
expertise_areas, created_at) and partitions every expert into groups.

Sizing rule
-----------
    ideal group size = 6,  min = 3,  max = 7
    * <= 7 experts                -> a single group
    * otherwise                   -> as many groups of ~6 as needed, sizes kept
                                     inside [3, 7] and balanced as evenly as
                                     possible. A trailing 1-2 "remainder" experts
                                     are absorbed into existing groups (pushing
                                     them to 7) rather than forming a sub-min
                                     group of 1 or 2.

Diversity heuristic
-------------------
Groups are filled to maximise *role* variety: each expert is placed in the
still-open group that currently holds the fewest members of that expert's role
(ties broken by most remaining capacity). Roles are dealt rare-first so uncommon
roles get spread across groups instead of clumping. Deterministic.

Outputs
-------
    group_output/expert_groups.csv   personal_id, affiliation, role, group_id (1..n)
    group_output/expert_groups.pptx  one "Groups" slide: a grid of Group ID cards

Usage
-----
    python group_experts.py
    python group_experts.py --input experts.csv --group-size 6
    python group_experts.py --no-slides
"""

from __future__ import annotations

import argparse
import math
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional, Sequence

import pandas as pd

HERE = Path(__file__).resolve().parent
DEFAULT_INPUT = HERE / "experts.csv"

UNSPEC = "Unspecified"


# --------------------------------------------------------------------------- #
# Data model
# --------------------------------------------------------------------------- #
@dataclass(frozen=True)
class Expert:
    personal_id: str
    affiliation: str
    role: str


def _clean(value: object) -> str:
    if value is None:
        return ""
    s = str(value).strip()
    if s.lower() in ("", "nan", "none"):
        return ""
    return s


def load_experts(path: Path) -> List[Expert]:
    df = pd.read_csv(path, dtype=str, keep_default_na=False)
    needed = {"personal_id", "affiliation", "role"}
    missing = needed - set(df.columns)
    if missing:
        raise SystemExit(f"Input is missing required columns: {sorted(missing)}")

    experts: List[Expert] = []
    for idx, row in df.iterrows():
        pid = _clean(row.get("personal_id")) or f"expert_{idx + 1}"
        experts.append(
            Expert(
                personal_id=pid,
                affiliation=_clean(row.get("affiliation")) or UNSPEC,
                role=_clean(row.get("role")) or UNSPEC,
            )
        )
    return experts


# --------------------------------------------------------------------------- #
# Group sizing
# --------------------------------------------------------------------------- #
def compute_group_sizes(n: int, ideal: int = 6, min_size: int = 3, max_size: int = 7) -> List[int]:
    """Return a list of group sizes summing to ``n`` and (where possible) inside
    [min_size, max_size], balanced as evenly as possible."""
    if n <= 0:
        return []
    if n <= max_size:
        return [n]  # one group, even if below min_size

    k = max(1, round(n / ideal))
    while k * max_size < n:            # too few groups -> some group would exceed max
        k += 1
    while k > 1 and math.ceil(n / k) < min_size:   # too many -> a group falls below min
        k -= 1

    base, rem = divmod(n, k)
    # `rem` groups get one extra member; this is exactly the "absorb the 1-2
    # leftover into existing groups" behaviour (e.g. 14 -> [7,7], 13 -> [7,6]).
    sizes = [base + 1] * rem + [base] * (k - rem)
    return sizes


# --------------------------------------------------------------------------- #
# Diversity assignment
# --------------------------------------------------------------------------- #
def assign_groups(experts: Sequence[Expert], sizes: Sequence[int]) -> List[List[Expert]]:
    """Greedily place experts so each group gets a varied mix of roles."""
    k = len(sizes)
    if k == 0:
        return []
    if k == 1:
        return [list(experts)]

    # Bucket by role; deal rare roles first so they spread across groups.
    buckets: Dict[str, List[Expert]] = defaultdict(list)
    for e in experts:
        buckets[e.role].append(e)
    ordered_roles = sorted(buckets, key=lambda r: (len(buckets[r]), r))

    # Interleave into one sequence so consecutive experts differ in role.
    sequence: List[Expert] = []
    pools = [list(buckets[r]) for r in ordered_roles]
    while any(pools):
        for pool in pools:
            if pool:
                sequence.append(pool.pop(0))

    groups: List[List[Expert]] = [[] for _ in range(k)]

    def role_count(group: List[Expert], role: str) -> int:
        return sum(1 for m in group if m.role == role)

    for e in sequence:
        best_gi, best_key = None, None
        for gi in range(k):
            if len(groups[gi]) >= sizes[gi]:
                continue
            # fewest of this role, then most remaining capacity, then lowest index
            key = (role_count(groups[gi], e.role), -(sizes[gi] - len(groups[gi])), gi)
            if best_key is None or key < best_key:
                best_key, best_gi = key, gi
        groups[best_gi].append(e)

    return groups


def build_groups(experts: Sequence[Expert], ideal: int = 6) -> List[List[Expert]]:
    sizes = compute_group_sizes(len(experts), ideal=ideal)
    return assign_groups(experts, sizes)


# --------------------------------------------------------------------------- #
# CSV output
# --------------------------------------------------------------------------- #
def groups_to_frame(groups: Sequence[Sequence[Expert]]) -> pd.DataFrame:
    rows = []
    for gid, group in enumerate(groups, start=1):
        for e in group:
            rows.append(
                dict(personal_id=e.personal_id, affiliation=e.affiliation,
                     role=e.role, group_id=gid)
            )
    return pd.DataFrame(rows, columns=["personal_id", "affiliation", "role", "group_id"])


# --------------------------------------------------------------------------- #
# Slide — "Groups" grid of Group ID cards (matches the mock)
# --------------------------------------------------------------------------- #
def build_slide(groups: Sequence[Sequence[Expert]], path: str) -> Optional[str]:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        from pptx.enum.shapes import MSO_SHAPE
    except ImportError:
        print("(python-pptx not installed -> skipping slide; pip install python-pptx)",
              file=sys.stderr)
        return None

    NAVY = RGBColor(0x15, 0x26, 0x4A)
    CARD = RGBColor(0xDD, 0xF0, 0xF3)      # light teal card
    PILL = RGBColor(0xFF, 0xFF, 0xFF)      # white label pill
    IDCOL = RGBColor(0x33, 0x3A, 0x45)     # member id text
    NUMCOL = RGBColor(0x9A, 0xA0, 0xAA)    # member number
    GRAY = RGBColor(0x55, 0x5E, 0x6B)
    SERIF, SANS = "Georgia", "Calibri"

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    def card(x, y, w, h, fill, radius=0.06):
        b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
        b.fill.solid(); b.fill.fore_color.rgb = fill
        b.line.fill.background(); b.shadow.inherit = False
        try:
            b.adjustments[0] = radius
        except Exception:
            pass
        return b

    def txt(x, y, w, h, anchor=MSO_ANCHOR.TOP):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True; tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.08)
        tf.margin_top = tf.margin_bottom = Inches(0.03)
        return tf

    def para(tf, runs, first=False, align=PP_ALIGN.LEFT, sb=0, sa=0, line=1.0):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        p.alignment = align; p.space_before = Pt(sb); p.space_after = Pt(sa)
        p.line_spacing = line
        for text, size, color, bold, font in runs:
            r = p.add_run(); r.text = text
            r.font.size = Pt(size); r.font.color.rgb = color
            r.font.bold = bold; r.font.name = font
        return p

    # Title
    para(txt(0.55, 0.30, 12.0, 0.8), [("Groups", 34, NAVY, True, SERIF)], first=True)

    # Grid geometry — fixed 4-column grid (cards stay compact, like the mock)
    cols = 4
    margin_x, gap_x, gap_y = 0.55, 0.30, 0.35
    top = 1.45
    card_w = (13.333 - 2 * margin_x - (cols - 1) * gap_x) / cols

    max_members = max((len(g) for g in groups), default=0)
    row_h = 0.235
    header_h = 0.62
    pad_bottom = 0.22
    card_h = header_h + max_members * row_h + pad_bottom

    pill_w, pill_h = 1.85, 0.42

    for i, group in enumerate(groups):
        col, row = i % cols, i // cols
        cx = margin_x + col * (card_w + gap_x)
        cy = top + row * (card_h + gap_y)

        card(cx, cy, card_w, card_h, CARD)

        # white "Group ID: N" pill, centred at the top of the card
        px = cx + (card_w - pill_w) / 2
        py = cy + 0.12
        card(px, py, pill_w, pill_h, PILL, radius=0.12)
        ptf = txt(px, py, pill_w, pill_h, MSO_ANCHOR.MIDDLE)
        para(ptf, [(f"Group ID: {i + 1}", 12.5, NAVY, True, SANS)],
             first=True, align=PP_ALIGN.CENTER)

        # numbered member list
        body = txt(cx + 0.22, cy + header_h + 0.02, card_w - 0.34,
                   card_h - header_h - 0.1)
        for j, e in enumerate(group, start=1):
            para(body, [(f"{j}.  ", 11, NUMCOL, False, SANS),
                        (e.personal_id, 12, IDCOL, True, SANS)],
                 first=(j == 1), sb=1, sa=1, line=1.0)

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #
def parse_args(argv: Optional[Sequence[str]] = None) -> argparse.Namespace:
    ap = argparse.ArgumentParser(description="Divide experts into role-diverse groups.")
    ap.add_argument("--input", type=Path, default=DEFAULT_INPUT,
                    help="path to experts.csv (default: ./experts.csv)")
    ap.add_argument("--group-size", type=int, default=6,
                    help="ideal group size (default: 6; min 3, max 7 enforced)")
    ap.add_argument("--out", type=Path, default=HERE / "group_output",
                    help="output directory")
    ap.add_argument("--csv-out", type=Path, default=None,
                    help="CSV path (default: <out>/expert_groups.csv)")
    ap.add_argument("--pptx", type=Path, default=None,
                    help="slide path (default: <out>/expert_groups.pptx)")
    ap.add_argument("--no-slides", action="store_true", help="skip slide generation")
    return ap.parse_args(argv)


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = parse_args(argv)
    if args.group_size <= 0:
        raise SystemExit("--group-size must be positive")
    if not args.input.is_file():
        raise SystemExit(f"Input file not found: {args.input}")

    experts = load_experts(args.input)
    if not experts:
        raise SystemExit(f"No experts loaded from {args.input}")

    groups = build_groups(experts, ideal=args.group_size)
    frame = groups_to_frame(groups)

    outdir = args.out
    outdir.mkdir(parents=True, exist_ok=True)
    csv_path = args.csv_out or (outdir / "expert_groups.csv")
    frame.to_csv(csv_path, index=False)

    # Console summary
    print(f"Loaded {len(experts)} experts from {args.input}")
    print(f"Created {len(groups)} group(s)  (sizes: {[len(g) for g in groups]})")
    for gid, group in enumerate(groups, start=1):
        roles = Counter(e.role for e in group)
        role_summary = ", ".join(f"{r} x{c}" for r, c in sorted(roles.items()))
        print(f"  Group {gid} ({len(group)}): "
              f"{', '.join(e.personal_id for e in group)}   [{role_summary}]")
    print(f"Wrote {csv_path}")

    if not args.no_slides:
        path = build_slide(groups, str(args.pptx or (outdir / "expert_groups.pptx")))
        if path:
            print(f"Wrote slide -> {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())